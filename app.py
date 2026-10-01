"""
Pastoral Comercial - Sistema de Metas de Recaudación
Flask backend con integración Excel
"""

import os
import json
from datetime import datetime, date, timedelta
from functools import wraps

import pandas as pd
from flask import (
    Flask, render_template, request, redirect, url_for,
    session, flash, jsonify
)
from werkzeug.security import generate_password_hash, check_password_hash
import openpyxl
from openpyxl import load_workbook, Workbook

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "pastoral-comercial-2026-secret-key")

# En producción (Railway) se puede montar un volumen en /data
# En local usa la carpeta del proyecto
DATA_DIR = os.environ.get("DATA_DIR", os.path.dirname(os.path.abspath(__file__)))
EXCEL_FILE = os.path.join(DATA_DIR, "datos.xlsx")

# ──────────────────────────────────────────────────────────────────────
# UTILIDADES EXCEL
# ──────────────────────────────────────────────────────────────────────

def init_excel():
    """Crea el archivo Excel con hojas y datos iniciales si no existe."""
    if os.path.exists(EXCEL_FILE):
        return
    wb = Workbook()

    # Hoja Usuarios
    ws_users = wb.active
    ws_users.title = "Usuarios"
    ws_users.append(["usuario", "nombre", "password_hash", "rol", "activo"])
    # Admin por defecto: usuario=admin, contraseña=admin123
    ws_users.append([
        "admin",
        "Administrador",
        generate_password_hash("admin123"),
        "admin",
        True
    ])
    ws_users.append([
        "juan",
        "Juan Pérez",
        generate_password_hash("juan123"),
        "usuario",
        True
    ])
    ws_users.append([
        "maria",
        "María García",
        generate_password_hash("maria123"),
        "usuario",
        True
    ])

    # Hoja Semanas (metas semanales)
    ws_weeks = wb.create_sheet("Semanas")
    ws_weeks.append(["id", "titulo", "descripcion", "meta", "fecha_inicio", "fecha_fin", "activa"])
    # Semana de ejemplo
    today = date.today()
    monday = today - timedelta(days=today.weekday())
    sunday = monday + timedelta(days=6)
    ws_weeks.append([
        1,
        "Semana 1 - Pastoral",
        "Primera semana de recaudación",
        500000,
        monday.strftime("%Y-%m-%d"),
        sunday.strftime("%Y-%m-%d"),
        True
    ])

    # Hoja Aportes
    ws_contributions = wb.create_sheet("Aportes")
    ws_contributions.append(["id", "usuario", "semana_id", "monto", "descripcion", "fecha"])
    ws_contributions.append([1, "juan", 1, 50000, "Donación inicial", today.strftime("%Y-%m-%d")])
    ws_contributions.append([2, "maria", 1, 75000, "Aporte personal", today.strftime("%Y-%m-%d")])

    wb.save(EXCEL_FILE)


def read_sheet(sheet_name):
    """Lee una hoja del Excel y retorna lista de dicts."""
    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=sheet_name, dtype=str)
        return df.where(pd.notna(df), None).to_dict(orient="records")
    except Exception:
        return []


def write_sheet(sheet_name, data, columns):
    """Escribe una lista de dicts en una hoja del Excel."""
    wb = load_workbook(EXCEL_FILE)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.append(columns)
    for row in data:
        ws.append([row.get(c) for c in columns])
    # Reordenar hojas para mantener el orden original
    sheet_order = ["Usuarios", "Semanas", "Aportes"]
    existing = [s for s in sheet_order if s in wb.sheetnames]
    others = [s for s in wb.sheetnames if s not in sheet_order]
    wb._sheets = [wb[s] for s in existing + others]
    wb.save(EXCEL_FILE)


def get_next_id(records):
    """Obtiene el siguiente ID para un registro."""
    if not records:
        return 1
    ids = []
    for r in records:
        try:
            ids.append(int(r.get("id") or 0))
        except (ValueError, TypeError):
            pass
    return max(ids) + 1 if ids else 1


# ──────────────────────────────────────────────────────────────────────
# FUNCIONES DE DATOS
# ──────────────────────────────────────────────────────────────────────

def get_users():
    return read_sheet("Usuarios")


def get_semanas():
    rows = read_sheet("Semanas")
    result = []
    for r in rows:
        try:
            r["id"] = int(r["id"]) if r.get("id") else None
            r["meta"] = float(r["meta"]) if r.get("meta") else 0
            r["activa"] = str(r.get("activa", "")).strip().lower() in ("true", "1", "yes", "sí", "si")
        except Exception:
            pass
        result.append(r)
    return result


def get_aportes():
    rows = read_sheet("Aportes")
    result = []
    for r in rows:
        try:
            r["id"] = int(r["id"]) if r.get("id") else None
            r["semana_id"] = int(r["semana_id"]) if r.get("semana_id") else None
            r["monto"] = float(r["monto"]) if r.get("monto") else 0
        except Exception:
            pass
        result.append(r)
    return result


def get_active_semana():
    for s in get_semanas():
        if s.get("activa"):
            return s
    semanas = get_semanas()
    return semanas[-1] if semanas else None


def get_ranking(semana_id=None):
    """Retorna ranking de usuarios con sus totales en la semana dada."""
    aportes = get_aportes()
    users = get_users()

    # Filtrar por semana si se especifica
    if semana_id is not None:
        aportes = [a for a in aportes if a.get("semana_id") == semana_id]

    # Sumar por usuario
    totales = {}
    for a in aportes:
        u = a.get("usuario")
        if u:
            totales[u] = totales.get(u, 0) + (a.get("monto") or 0)

    # Construir ranking con info de usuario
    ranking = []
    for user in users:
        if user.get("rol") == "admin" or str(user.get("activo", "")).lower() not in ("true", "1"):
            continue
        u = user.get("usuario")
        ranking.append({
            "usuario": u,
            "nombre": user.get("nombre", u),
            "total": totales.get(u, 0),
        })

    ranking.sort(key=lambda x: x["total"], reverse=True)
    for i, r in enumerate(ranking):
        r["posicion"] = i + 1

    return ranking


def get_user_aportes(username, semana_id=None):
    aportes = get_aportes()
    result = [a for a in aportes if a.get("usuario") == username]
    if semana_id:
        result = [a for a in result if a.get("semana_id") == semana_id]
    return sorted(result, key=lambda x: x.get("fecha") or "", reverse=True)


# ──────────────────────────────────────────────────────────────────────
# AUTENTICACIÓN
# ──────────────────────────────────────────────────────────────────────

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "usuario" not in session:
            flash("Debés iniciar sesión primero.", "warning")
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated


def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "usuario" not in session:
            return redirect(url_for("login"))
        if session.get("rol") != "admin":
            flash("Acceso restringido a administradores.", "danger")
            return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return decorated


# ──────────────────────────────────────────────────────────────────────
# RUTAS - AUTH
# ──────────────────────────────────────────────────────────────────────

@app.route("/", methods=["GET", "POST"])
def login():
    if "usuario" in session:
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        username = request.form.get("usuario", "").strip()
        password = request.form.get("password", "")

        users = get_users()
        user = next((u for u in users if u.get("usuario") == username), None)

        if user and check_password_hash(user["password_hash"], password):
            if str(user.get("activo", "")).lower() not in ("true", "1"):
                flash("Tu cuenta está desactivada.", "danger")
                return redirect(url_for("login"))
            session["usuario"] = username
            session["nombre"] = user.get("nombre", username)
            session["rol"] = user.get("rol", "usuario")
            return redirect(url_for("dashboard"))
        else:
            flash("Usuario o contraseña incorrectos.", "danger")

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ──────────────────────────────────────────────────────────────────────
# RUTAS - DASHBOARD USUARIO
# ──────────────────────────────────────────────────────────────────────

@app.route("/dashboard")
@login_required
def dashboard():
    semana_activa = get_active_semana()
    semana_id = semana_activa["id"] if semana_activa else None
    ranking = get_ranking(semana_id)
    user_aportes = get_user_aportes(session["usuario"], semana_id)
    user_total = sum(a.get("monto", 0) for a in user_aportes)
    user_pos = next((r["posicion"] for r in ranking if r["usuario"] == session["usuario"]), None)
    meta = float(semana_activa["meta"]) if semana_activa else 0
    progreso = min((user_total / meta * 100) if meta > 0 else 0, 100)
    total_global = sum(r["total"] for r in ranking)
    progreso_global = min((total_global / meta * 100) if meta > 0 else 0, 100)

    return render_template(
        "dashboard.html",
        semana=semana_activa,
        ranking=ranking,
        user_aportes=user_aportes,
        user_total=user_total,
        user_pos=user_pos,
        meta=meta,
        progreso=round(progreso, 1),
        total_global=total_global,
        progreso_global=round(progreso_global, 1),
    )


@app.route("/ranking")
@login_required
def ranking_page():
    semanas = get_semanas()
    semana_id_sel = request.args.get("semana_id")

    if semana_id_sel:
        try:
            semana_id_sel = int(semana_id_sel)
        except ValueError:
            semana_id_sel = None

    if semana_id_sel is None:
        active = get_active_semana()
        semana_id_sel = active["id"] if active else None

    semana_actual = next((s for s in semanas if s.get("id") == semana_id_sel), None)
    ranking = get_ranking(semana_id_sel)
    meta = float(semana_actual["meta"]) if semana_actual else 0
    total_global = sum(r["total"] for r in ranking)

    return render_template(
        "ranking.html",
        ranking=ranking,
        semanas=semanas,
        semana_actual=semana_actual,
        semana_id_sel=semana_id_sel,
        meta=meta,
        total_global=total_global,
    )


@app.route("/mis-aportes")
@login_required
def mis_aportes():
    semanas = get_semanas()
    all_aportes = get_user_aportes(session["usuario"])
    # Calcular total por semana
    totales = {}
    for a in all_aportes:
        sid = a.get("semana_id")
        totales[sid] = totales.get(sid, 0) + a.get("monto", 0)

    return render_template(
        "mis_aportes.html",
        aportes=all_aportes,
        semanas={s["id"]: s for s in semanas},
        totales=totales,
    )


# ──────────────────────────────────────────────────────────────────────
# RUTAS - ADMIN
# ──────────────────────────────────────────────────────────────────────

@app.route("/admin")
@admin_required
def admin_panel():
    users = [u for u in get_users() if u.get("rol") != "admin"]
    semanas = get_semanas()
    aportes = get_aportes()
    semana_activa = get_active_semana()
    ranking = get_ranking(semana_activa["id"] if semana_activa else None)
    total_recaudado = sum(r["total"] for r in ranking)
    return render_template(
        "admin.html",
        users=users,
        semanas=semanas,
        aportes=aportes,
        semana_activa=semana_activa,
        ranking=ranking,
        total_recaudado=total_recaudado,
    )


# ── Admin: Gestión de usuarios ──

@app.route("/admin/usuarios/crear", methods=["POST"])
@admin_required
def crear_usuario():
    users = get_users()
    username = request.form.get("usuario", "").strip().lower()
    nombre = request.form.get("nombre", "").strip()
    password = request.form.get("password", "")
    rol = request.form.get("rol", "usuario")

    if not username or not nombre or not password:
        flash("Todos los campos son obligatorios.", "danger")
        return redirect(url_for("admin_panel"))

    if any(u.get("usuario") == username for u in users):
        flash(f"El usuario '{username}' ya existe.", "danger")
        return redirect(url_for("admin_panel"))

    users.append({
        "usuario": username,
        "nombre": nombre,
        "password_hash": generate_password_hash(password),
        "rol": rol,
        "activo": True,
    })
    write_sheet("Usuarios", users, ["usuario", "nombre", "password_hash", "rol", "activo"])
    flash(f"Usuario '{nombre}' creado exitosamente.", "success")
    return redirect(url_for("admin_panel"))


@app.route("/admin/usuarios/toggle/<username>")
@admin_required
def toggle_usuario(username):
    users = get_users()
    for u in users:
        if u.get("usuario") == username:
            current = str(u.get("activo", "")).lower() in ("true", "1")
            u["activo"] = not current
            break
    write_sheet("Usuarios", users, ["usuario", "nombre", "password_hash", "rol", "activo"])
    flash("Estado del usuario actualizado.", "success")
    return redirect(url_for("admin_panel"))


@app.route("/admin/usuarios/eliminar/<username>")
@admin_required
def eliminar_usuario(username):
    users = [u for u in get_users() if u.get("usuario") != username]
    write_sheet("Usuarios", users, ["usuario", "nombre", "password_hash", "rol", "activo"])
    flash("Usuario eliminado.", "success")
    return redirect(url_for("admin_panel"))


# ── Admin: Gestión de semanas ──

@app.route("/admin/semanas/crear", methods=["POST"])
@admin_required
def crear_semana():
    semanas = get_semanas()
    titulo = request.form.get("titulo", "").strip()
    descripcion = request.form.get("descripcion", "").strip()
    meta = request.form.get("meta", "0")
    fecha_inicio = request.form.get("fecha_inicio", "")
    fecha_fin = request.form.get("fecha_fin", "")
    activar = request.form.get("activar") == "on"

    if not titulo or not meta:
        flash("Título y meta son obligatorios.", "danger")
        return redirect(url_for("admin_panel"))

    # Si se activa esta, desactivar las demás
    if activar:
        for s in semanas:
            s["activa"] = False

    new_id = get_next_id(semanas)
    semanas.append({
        "id": new_id,
        "titulo": titulo,
        "descripcion": descripcion,
        "meta": meta,
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "activa": activar,
    })
    write_sheet("Semanas", semanas, ["id", "titulo", "descripcion", "meta", "fecha_inicio", "fecha_fin", "activa"])
    flash(f"Semana '{titulo}' creada exitosamente.", "success")
    return redirect(url_for("admin_panel"))


@app.route("/admin/semanas/activar/<int:semana_id>")
@admin_required
def activar_semana(semana_id):
    semanas = get_semanas()
    for s in semanas:
        s["activa"] = (s.get("id") == semana_id)
    write_sheet("Semanas", semanas, ["id", "titulo", "descripcion", "meta", "fecha_inicio", "fecha_fin", "activa"])
    flash("Semana activada.", "success")
    return redirect(url_for("admin_panel"))


@app.route("/admin/semanas/eliminar/<int:semana_id>")
@admin_required
def eliminar_semana(semana_id):
    semanas = [s for s in get_semanas() if s.get("id") != semana_id]
    write_sheet("Semanas", semanas, ["id", "titulo", "descripcion", "meta", "fecha_inicio", "fecha_fin", "activa"])
    flash("Semana eliminada.", "success")
    return redirect(url_for("admin_panel"))


# ── Admin: Gestión de aportes ──

@app.route("/admin/aportes/registrar", methods=["POST"])
@admin_required
def registrar_aporte():
    aportes = get_aportes()
    usuario = request.form.get("usuario", "").strip()
    semana_id = request.form.get("semana_id", "")
    monto = request.form.get("monto", "0")
    descripcion = request.form.get("descripcion", "").strip()
    fecha = request.form.get("fecha") or date.today().strftime("%Y-%m-%d")

    if not usuario or not semana_id or not monto:
        flash("Usuario, semana y monto son obligatorios.", "danger")
        return redirect(url_for("admin_panel"))

    new_id = get_next_id(aportes)
    aportes.append({
        "id": new_id,
        "usuario": usuario,
        "semana_id": int(semana_id),
        "monto": float(monto.replace(",", ".")),
        "descripcion": descripcion,
        "fecha": fecha,
    })
    write_sheet("Aportes", aportes, ["id", "usuario", "semana_id", "monto", "descripcion", "fecha"])
    flash("Aporte registrado exitosamente.", "success")
    return redirect(url_for("admin_panel"))


@app.route("/admin/aportes/eliminar/<int:aporte_id>")
@admin_required
def eliminar_aporte(aporte_id):
    aportes = [a for a in get_aportes() if a.get("id") != aporte_id]
    write_sheet("Aportes", aportes, ["id", "usuario", "semana_id", "monto", "descripcion", "fecha"])
    flash("Aporte eliminado.", "success")
    return redirect(url_for("admin_panel"))


# ── API JSON para gráficos ──

@app.route("/api/ranking")
@login_required
def api_ranking():
    semana_id = request.args.get("semana_id")
    if semana_id:
        try:
            semana_id = int(semana_id)
        except ValueError:
            semana_id = None
    ranking = get_ranking(semana_id)
    return jsonify(ranking)


# ──────────────────────────────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    init_excel()
    import socket
    try:
        hostname = socket.gethostname()
        local_ip = socket.gethostbyname(hostname)
    except Exception:
        local_ip = "tu-ip-local"

    print("\n" + "="*60)
    print("  PASTORAL INGENIERIA COMERCIAL UC")
    print("="*60)
    print("  Red local (WiFi mismo edificio):")
    print(f"    http://{local_ip}:5000")
    print("  Solo esta PC:")
    print("    http://localhost:5000")
    print("")
    print("  Admin: usuario=admin  contrasena=admin123")
    print("="*60 + "\n")
    app.run(host="0.0.0.0", port=5000, debug=False)
